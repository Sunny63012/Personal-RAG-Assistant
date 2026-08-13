import fitz
import io
import pandas as pd
import streamlit as st
from langchain_core.documents import Document
from docx import Document as DocxDocument
# def load_documents(file_paths):
#     documents = []
#     for file_path in file_paths:
#         #PDF
#         if file_path.endswith(".pdf"):
#             doc = fitz.open(file_path)#doc is a PyMuPDF Document object that represents the opened PDF file. It allows you to access the pages and their content.
#             text = ""
#             for page in doc:
#                 text = page.get_text()
#                 documents.append(
#                     Document(
#                     page_content=text,
#                     metadata={
#                         "filename":os.path.basename(file_path),
#                         "page_number": page.number + 1,
#                         "source_type":"pdf"
#                     })
#                 )
#         #Excel
#         elif file_path.endswith((".xlsx", ".xls")):
#             df = pd.read_excel(file_path)
#             text = df.to_string(index=False)#convert the DataFrame to a string representation without the index
#             documents.append(Document(
#                 page_content=text,
#                 metadata={
#                     "filename": os.path.basename(file_path),
#                     "source_type": "excel"
#                 }))
#         #Word
#         elif file_path.endswith(".docx"):
#             doc = DocxDocument(file_path)
#             text = "\n".join([para.text for para in doc.paragraphs])#we use a list comprehension to iterate over each paragraph in the document and extract its text. 
#             #he resulting list of paragraph texts is then joined together into a single string, with each paragraph separated by a newline character (\n). This creates a coherent representation of the entire document's content.
#             documents.append(Document(
#                 page_content=text,
#                 metadata={
#                     "filename": os.path.basename(file_path),
#                     "source_type": "docx"
#                 }))
#         #Text
#         else:
#             with open(file_path, "r",encoding="utf-8") as f:
#                 text=f.read()
#                 documents.append(Document(
#                     page_content=text,
#                     metadata={
#                         "filename": os.path.basename(file_path),
#                         "source_type": "txt"
#                     }))
#     return documents
def load_documents(uploaded_files):
    documents = []
    tables={}
    for uploaded_file in uploaded_files:
        name = uploaded_file.name
        try:
            if name.lower().endswith(".pdf"):
                documents.extend(_load_pdf(uploaded_file, name))
            elif name.lower().endswith((".xlsx", ".xls")):
                docs,sheets=_load_excel(uploaded_file, name)
                documents.extend(docs)
                tables.update(sheets)
            elif name.lower().endswith(".docx"):
                documents.extend(_load_docx(uploaded_file, name))
            elif name.lower().endswith(".csv"):
                docs,df=_load_csv(uploaded_file, name)
                documents.extend(docs)
                if df is not None:
                    tables[name]=df
            else:
                documents.extend(_load_text(uploaded_file, name))
        except Exception as e:
            # One corrupt/unsupported file should never take down the whole batch.
            st.warning(f"Skipped '{name}' — couldn't read it ({e}).")
    return documents,tables
def _load_pdf(uploaded_file,name):
    docs=[]
    pdf_bytes=uploaded_file.read()
    doc=fitz.open(stream=pdf_bytes,filetype="pdf")
    for page in doc:
        text = page.get_text()
        if not text.strip():
            # Empty text almost always means a scanned/image-only page.
            # We don't OCR here, but we tell the user instead of silently
            # dropping content they'll expect to be searchable.
            st.warning(f"Page {page.number + 1} of '{name}' looks scanned — no text extracted.")
            continue
        docs.append(Document(
            page_content=text,
            metadata={"filename": name, "page_number": page.number + 1, "source_type": "pdf"}
        ))
    return docs
def _load_excel(uploaded_file,name):
    docs=[]
    tables={}
    excel_bytes=uploaded_file.read()
     # sheet_name=None loads every sheet into a dict {sheet_name: DataFrame}
    sheets = pd.read_excel(io.BytesIO(excel_bytes), sheet_name=None)
    for sheet_name, df in sheets.items():
        if df.empty:
            continue
        tables[f"{name} [{sheet_name}]"] = df
        for row_number,row in df.iterrows():
            row_data=[]
            for column in df.columns:
                value=row[column]
                if pd.notna(value):
                    row_data.append(f"{column}:{value}")
            text="\n".join(row_data)
            docs.append(Document(
                page_content=text,
                metadata={"filename": name, "sheet_name": sheet_name,"row_number":row_number+2, "source_type": "excel"}
            ))
    return docs,tables
def _load_csv(uploaded_file, name):
    docs=[]
    csv_bytes = uploaded_file.read()
    df = pd.read_csv(io.BytesIO(csv_bytes))
    if df.empty:
        return []
    for row_number, row in df.iterrows():
        row_data=[]
        for column in df.columns:
            value=row[column]
            if pd.notna(value):
                row_data.append(f"{column}:{value}")
        text="\n".join(row_data)
        docs.append(Document(
            page_content=text,
            metadata={"filename": name, "row_number": row_number + 2, "source_type": "csv"}
        ))
    return docs,df
def _load_docx(uploaded_file, name):
    docx_bytes = uploaded_file.read()
    doc = DocxDocument(io.BytesIO(docx_bytes))

    parts = [para.text for para in doc.paragraphs if para.text.strip()]

    # Tables are where most of the real data lives in business docs —
    # pull them out as pipe-separated rows so the splitter can chunk them sanely.
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    text = "\n".join(parts)
    return [Document(
        page_content=text,
        metadata={"filename": name, "source_type": "docx"}
    )]
def _load_text(uploaded_file, name):
    raw = uploaded_file.read()
    # errors="replace" instead of raising on non-UTF-8 files
    text = raw.decode("utf-8", errors="replace")
    return [Document(
        page_content=text,
        metadata={"filename": name, "source_type": "txt"}
    )]

    